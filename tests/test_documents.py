import shutil
from pathlib import Path

import numpy as np
import pandas as pd
import pytest
from docx import Document
from matplotlib import pyplot as plt
from pptx import Presentation

from agentic_conservation.documents import (
    FIGURE,
    MAIN_FIGURES,
    MAIN_TABLES,
    TABLE,
    _figure_deck,
    _manuscript_word_count,
    _parenthetical_citations,
    _reference_text,
    _validate_literature,
    _validate_manuscript_word_limit,
    _validate_references,
)


def _write_counting_fixture(path: Path) -> None:
    document = Document()
    document.add_paragraph("Excluded title words")
    document.add_heading("Abstract", level=1)
    document.add_paragraph("One two three")
    document.add_paragraph("Keywords: four; five")
    document.add_heading("Methods", level=1)
    document.add_paragraph("Six seven")
    document.add_paragraph(f"Figure 1. {MAIN_FIGURES[0][2]}")
    document.add_paragraph("Figure 1. Narrative starts here")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Excluded table words"
    document.add_heading("Acknowledgments", level=1)
    document.add_paragraph("Eight nine")
    document.add_heading("Data and code availability", level=1)
    document.add_paragraph("Excluded trailing words")
    document.save(path)


def test_manuscript_word_count_uses_journal_range_and_exclusions(tmp_path):
    path = tmp_path / "manuscript.docx"
    _write_counting_fixture(path)
    assert _manuscript_word_count(path) == 18


def test_literature_uses_author_year_citations_and_unnumbered_references():
    literature = pd.DataFrame(
        [
            {
                "reference_authors": "Recent AB",
                "citation": "Recent 2020",
                "year": 2020,
                "title": "Recent paper",
                "journal": "Complete Journal Name",
                "doi": "10.0/recent",
            },
            {
                "reference_authors": "Earlier CD",
                "citation": "Earlier 2010",
                "year": 2010,
                "title": "Earlier paper",
                "journal": "Complete Journal Name",
                "doi": "10.0/earlier",
            },
        ]
    )

    assert _parenthetical_citations(literature) == "(Earlier 2010; Recent 2020)"
    references = _reference_text(literature)
    assert references[0].startswith("Recent AB. 2020.")
    assert not references[0][0].isdigit()


def test_manuscript_word_limit_is_enforced(tmp_path):
    at_limit = tmp_path / "at-limit.docx"
    over_limit = tmp_path / "over-limit.docx"
    for path, word_count in [(at_limit, 6998), (over_limit, 6999)]:
        document = Document()
        document.add_heading("Abstract", level=1)
        document.add_paragraph("word " * word_count)
        document.add_heading("Acknowledgments", level=1)
        document.add_heading("Data and code availability", level=1)
        document.save(path)

    assert _validate_manuscript_word_limit(at_limit) == 7000
    with pytest.raises(ValueError, match="Manuscript is 7001 words.*limit is 7000"):
        _validate_manuscript_word_limit(over_limit)


def test_word_count_uses_acknowledgments_boundary_after_reordered_sections(tmp_path):
    path = tmp_path / "reordered.docx"
    document = Document()
    document.add_heading("Abstract", level=1)
    document.add_paragraph("word " * 6993)
    document.add_heading("Data and code availability", level=1)
    document.add_paragraph("included")
    document.add_heading("Acknowledgments", level=1)
    document.add_paragraph("included")
    document.add_heading("Literature cited", level=1)
    document.add_paragraph("excluded trailing words")
    document.save(path)

    assert _manuscript_word_count(path) == 7001
    with pytest.raises(ValueError, match="Manuscript is 7001 words.*limit is 7000"):
        _validate_manuscript_word_limit(path)


def test_word_count_includes_acknowledgments_subsections(tmp_path):
    path = tmp_path / "acknowledgments-subsection.docx"
    document = Document()
    document.add_heading("Abstract", level=1)
    document.add_paragraph("word " * 6997)
    document.add_heading("Acknowledgments", level=1)
    document.add_heading("Funding", level=2)
    document.add_paragraph("included")
    document.add_heading("Literature cited", level=1)
    document.add_paragraph("excluded trailing words")
    document.save(path)

    assert _manuscript_word_count(path) == 7001
    with pytest.raises(ValueError, match="Manuscript is 7001 words.*limit is 7000"):
        _validate_manuscript_word_limit(path)


def test_word_count_ignores_plain_acknowledgments_before_abstract(tmp_path):
    path = tmp_path / "plain-acknowledgments.docx"
    document = Document()
    document.add_paragraph("Acknowledgments")
    document.add_heading("Abstract", level=1)
    document.add_paragraph("word " * 6998)
    document.add_heading("Acknowledgments", level=1)
    document.add_paragraph("included")
    document.add_heading("Literature cited", level=1)
    document.save(path)

    assert _manuscript_word_count(path) == 7001
    with pytest.raises(ValueError, match="Manuscript is 7001 words.*limit is 7000"):
        _validate_manuscript_word_limit(path)


def test_figure_deck_uses_neutral_name_and_discloses_image_artwork(tmp_path):
    figures = tmp_path / "figures"
    manuscript = tmp_path / "manuscript"
    figures.mkdir()
    manuscript.mkdir()
    figure_specs = MAIN_FIGURES
    first = figures / figure_specs[0][1]
    plt.imsave(first, np.ones((2, 2, 3)))
    for _, filename, _ in figure_specs[1:]:
        shutil.copy2(first, figures / filename)
    legacy = manuscript / "EDITABLE_FIGURES.pptx"
    legacy.write_bytes(b"legacy")

    output = _figure_deck(tmp_path)
    presentation = Presentation(output)

    assert output.name == "FIGURE_DECK.pptx"
    assert not legacy.exists()
    assert presentation.core_properties.title == "Submission figure deck"
    assert "image-based figure artwork" in presentation.core_properties.subject
    assert len(presentation.slides) == len(figure_specs)
    for slide in presentation.slides:
        slide_text = " ".join(
            shape.text for shape in slide.shapes if hasattr(shape, "text_frame")
        )
        assert "Image-based artwork" in slide_text


def test_displays_use_one_main_text_sequence():
    assert [spec[0] for spec in MAIN_FIGURES] == [
        f"Figure {number}" for number in range(1, len(MAIN_FIGURES) + 1)
    ]
    assert [spec[0] for spec in MAIN_TABLES] == [
        f"Table {number}" for number in range(1, len(MAIN_TABLES) + 1)
    ]
    assert list(FIGURE.values()) == [spec[0] for spec in MAIN_FIGURES]
    assert list(TABLE.values()) == [spec[0] for spec in MAIN_TABLES]


def test_display_citations_must_precede_displays_in_order(tmp_path):
    path = tmp_path / "manuscript.docx"
    document = Document()
    document.add_paragraph("See Figure 1.")
    document.add_paragraph("Figure 1. Caption.")
    document.add_paragraph("See Figure 2.")
    document.add_paragraph("Figure 2. Caption.")
    document.save(path)
    _validate_references(path, ["Figure 1", "Figure 2"], [])
    with pytest.raises(ValueError):
        _validate_references(path, ["Figure 2", "Figure 1"], [])


def test_literature_validation_requires_cited_ordered_references(tmp_path):
    literature = pd.DataFrame(
        {"citation": ["Earlier 2010", "Recent 2020"], "year": [2010, 2020]}
    )
    cases = {
        "valid": ("Text (Earlier 2010; Recent 2020).", None),
        "uncited": ("Text (Earlier 2010).", "not cited"),
        "orphan": ("Text (Earlier 2010; Missing 2015; Recent 2020).", "without references"),
        "unordered": ("Text (Recent 2020; Earlier 2010).", "not chronological"),
    }
    for name, (body, error) in cases.items():
        path = tmp_path / f"{name}.docx"
        document = Document()
        document.add_paragraph(body)
        document.add_paragraph("Literature cited")
        document.add_paragraph("Earlier CD. 2010. Recent AB. 2020.")
        document.save(path)
        if error is None:
            _validate_literature(path, literature)
        else:
            with pytest.raises(ValueError, match=error):
                _validate_literature(path, literature)
