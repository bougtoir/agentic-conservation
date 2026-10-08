"""Generate anonymous Word and PowerPoint submission-support documents."""

from __future__ import annotations

import hashlib
import json
import math
import re
import struct
import zipfile
from collections.abc import Callable
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches as PptxInches
from pptx.util import Pt as PptxPt

_FIGURE_SPECS = [
    (
        "network",
        "fig1_conservation_network.png",
        "Conservation network represented in the simulation. Arrows denote "
        "potential movement among wild, managed, transitional, "
        "rehabilitation, and biobank nodes; feasible policies act on "
        "observed rather than true state."
    ),
    (
        "pareto",
        "fig3_pareto_outcomes.png",
        "Multiobjective policy comparison and Pareto status under the "
        "declared aggregate objectives."
    ),
    (
        "trajectories",
        "fig2_century_trajectories.png",
        "Mean 100-year trajectories under the eight conservation policies. "
        "Trajectories are synthetic model outcomes, not forecasts for real "
        "taxa."
    ),
    (
        "ebd",
        "fig5_extinction_before_discovery.png",
        "Extinction-before-discovery outcomes across latent-species "
        "fraction and survey-cost conditions."
    ),
    (
        "tradeoffs",
        "fig6_objective_tradeoffs.png",
        "Normalized cross-objective outcomes. Normalization is descriptive "
        "and does not define cardinal utility."
    ),
    (
        "observation",
        "fig7_partial_observation_architecture.png",
        "Partial-observation architecture separating true ecological "
        "state, observations, policy actions, and realized outcomes."
    ),
    (
        "reversal",
        "fig4_policy_reversal_thresholds.png",
        "Policy-reversal surfaces for selected two-factor sweeps. Color "
        "represents the difference in 100-year survival between the hybrid "
        "portfolio (S7) and habitat-first policy (S6)."
    ),
    (
        "exploration",
        "fig8_exploration_exploitation.png",
        "Exploration–exploitation patterns across survey-cost conditions."
    ),
    (
        "origin_future",
        "fig9_origin_future_optimum.png",
        "Origin-to-future habitat-quality shift under the deterministic "
        "component of directional environmental change."
    ),
]

_TABLE_SPECS = [
    (
        "parameters",
        "table1_model_parameters.csv",
        "Base simulation parameters."
    ),
    (
        "policies",
        "table2_policies.csv",
        "Policy definitions and budget shares."
    ),
    (
        "sensitivity",
        "table4_sensitivity.csv",
        "Design-level Latin-hypercube sensitivity correlations."
    ),
    (
        "interactions",
        "table4_sensitivity_interactions.csv",
        "Strongest pairwise partial-rank interaction screens."
    ),
    (
        "coverage",
        "table5_empirical_coverage.csv",
        "GBIF observation-coverage aggregates from persisted API "
        "responses."
    ),
    (
        "evidence",
        "table6_empirical_evidence.csv",
        "Bibliographic metadata and evidence-support levels for human "
        "appraisal."
    ),
    (
        "outcomes",
        "table3_policy_outcomes.csv",
        "Policy outcome estimates and replicate-distribution intervals."
    ),
    (
        "robustness",
        "table7_robustness.csv",
        "Ablation and adverse-condition outcomes."
    ),
    (
        "information_weights",
        "table8_information_weight_sensitivity.csv",
        "Current-state-information effect under alternative normative "
        "weights."
    ),
]

FIGURE = {key: f"Figure {index}" for index, (key, _, _) in enumerate(_FIGURE_SPECS, 1)}
TABLE = {key: f"Table {index}" for index, (key, _, _) in enumerate(_TABLE_SPECS, 1)}
MAIN_FIGURES = [(FIGURE[key], filename, caption) for key, filename, caption in _FIGURE_SPECS]
MAIN_TABLES = [(TABLE[key], filename, caption) for key, filename, caption in _TABLE_SPECS]
_FIGURE_BY_KEY = dict(zip(FIGURE, MAIN_FIGURES))
_TABLE_BY_KEY = dict(zip(TABLE, MAIN_TABLES))
COMPACT_TABLES = {
    TABLE[key]
    for key in [
        "sensitivity",
        "interactions",
        "coverage",
        "evidence",
        "robustness",
        "information_weights",
    ]
}


TITLE = "Dynamic conservation networks under environmental change and hidden biodiversity"

def _configure_document(document: Document, title: str) -> None:
    section = document.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    normal.paragraph_format.space_after = Pt(0)
    for style_name in ["Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"]:
        style = document.styles[style_name]
        style.font.name = "Times New Roman"
        style.font.color.rgb = RGBColor(0, 0, 0)
    document.core_properties.title = title
    document.core_properties.subject = "Anonymous human-review draft"
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    document.core_properties.comments = ""
    document.core_properties.keywords = ""
    document.core_properties.category = ""
    document.core_properties.revision = 1


def _add_title(document: Document, title: str, subtitle: str | None = None) -> None:
    paragraph = document.add_paragraph()
    paragraph.style = document.styles["Title"]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run(title)
    if subtitle:
        paragraph = document.add_paragraph()
        paragraph.style = document.styles["Subtitle"]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run(subtitle)


def _add_review_notice(document: Document) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(10)
    paragraph.paragraph_format.space_after = Pt(10)
    run = paragraph.add_run(
        "ANONYMOUS HUMAN-REVIEW DRAFT — NOT READY FOR SCIENTIFIC SUBMISSION"
    )
    run.bold = True
    run.font.color.rgb = RGBColor(160, 0, 0)
    paragraph = document.add_paragraph()
    paragraph.add_run(
        "All numerical results are generated from the current repository outputs. "
        "Complete text was appraised for the four case and decision-theory sources; "
        "other cited sources support only statements in their retained author "
        "abstracts or titles. Accountable human authors must verify all claims and "
        "assumptions, revise the prose in their own scientific voice, and approve "
        "the final submission."
    )


def _add_heading(document: Document, text: str, level: int = 1) -> None:
    paragraph = document.add_heading(text, level=level)
    paragraph.paragraph_format.keep_with_next = True


def _add_paragraph(document: Document, text: str, bold_prefix: str | None = None) -> None:
    paragraph = document.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        paragraph.add_run(bold_prefix).bold = True
        paragraph.add_run(text[len(bold_prefix) :])
    else:
        paragraph.add_run(text)


def _format_cell(value: object) -> str:
    if pd.isna(value):
        return "—"
    if isinstance(value, float):
        if math.isclose(value, round(value), abs_tol=1e-10):
            return str(int(round(value)))
        return f"{value:.3f}"
    return str(value)


def _shade_cell(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def _repeat_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def _set_table_font(table, size: float) -> None:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(size)


def _add_table(
    document: Document,
    frame: pd.DataFrame,
    number: str,
    caption: str,
    font_size: float = 8,
) -> None:
    caption_paragraph = document.add_paragraph()
    caption_paragraph.paragraph_format.space_before = Pt(14)
    caption_paragraph.paragraph_format.space_after = Pt(6)
    caption_paragraph.add_run(f"{number}. {caption}").bold = True
    table = document.add_table(rows=1, cols=len(frame.columns))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header = table.rows[0]
    _repeat_header(header)
    for index, column in enumerate(frame.columns):
        header.cells[index].text = str(column).replace("_", " ")
        _shade_cell(header.cells[index], "D9EAF7")
        for run in header.cells[index].paragraphs[0].runs:
            run.bold = True
    for row in frame.itertuples(index=False, name=None):
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = _format_cell(value)
    _set_table_font(table, font_size)
    document.add_paragraph()


def _add_display_table(
    document: Document,
    frame: pd.DataFrame,
    number: str,
    caption: str,
    font_size: float = 8,
) -> None:
    if len(frame.columns) <= 6:
        _add_table(document, frame, number, caption, font_size)
        return
    landscape = document.add_section(WD_SECTION.NEW_PAGE)
    landscape.orientation = WD_ORIENT.LANDSCAPE
    landscape.page_width = Inches(11)
    landscape.page_height = Inches(8.5)
    _add_table(document, frame, number, caption, font_size)
    portrait = document.add_section(WD_SECTION.NEW_PAGE)
    portrait.orientation = WD_ORIENT.PORTRAIT
    portrait.page_width = Inches(8.5)
    portrait.page_height = Inches(11)


def _add_figure(document: Document, path: Path, number: str, caption: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(14)
    paragraph.add_run().add_picture(str(path), width=Inches(6.2))
    caption_paragraph = document.add_paragraph()
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    caption_paragraph.paragraph_format.space_before = Pt(14)
    caption_paragraph.paragraph_format.space_after = Pt(8)
    run = caption_paragraph.add_run(f"{number}. {caption}")
    run.italic = True


def _main_figure(document: Document, root: Path, key: str) -> None:
    number, filename, caption = _FIGURE_BY_KEY[key]
    _add_figure(document, root / "figures" / filename, number, caption)


def _main_table(document: Document, root: Path, key: str) -> None:
    number, filename, caption = _TABLE_BY_KEY[key]
    _add_display_table(
        document,
        pd.read_csv(root / "tables" / filename),
        number,
        caption,
        font_size=7 if number in COMPACT_TABLES else 8,
    )


def _summary_row(summary: pd.DataFrame, policy: str, metric: str) -> pd.Series:
    return summary[
        (summary.policy == policy) & (summary.metric == metric) & (~summary.oracle)
    ].iloc[0]


def _interval(summary: pd.DataFrame, policy: str, metric: str) -> str:
    row = _summary_row(summary, policy, metric)
    return f"{row.estimate:.3f} [{row.lower:.3f}, {row.upper:.3f}]"


def _sweep_counts(root: Path) -> tuple[int, int, int]:
    summary = pd.read_csv(root / "results/threshold_sweep_audit.csv")
    return (
        int(summary.robust_s6_higher.sum()),
        int(summary.robust_s7_higher.sum()),
        int(summary.indeterminate.sum()),
    )


def _reference_text(literature: pd.DataFrame) -> list[str]:
    references = []
    for row in literature.to_dict("records"):
        volume = "" if pd.isna(row.get("volume")) else str(row["volume"])
        pages = "" if pd.isna(row.get("pages")) else str(row["pages"])
        locator = f" {volume}:{pages}" if volume and pages else f" {volume}" if volume else ""
        references.append(
            f"{row['reference_authors']}. {int(row['year'])}. {row['title']}. "
            f"{row['journal']}{locator}. DOI: {row['doi']}."
        )
    return references


def _parenthetical_citations(literature: pd.DataFrame) -> str:
    ordered = literature.sort_values(["year", "citation"])
    return "(" + "; ".join(ordered.citation) + ")"


def _citer(literature: pd.DataFrame) -> Callable[..., str]:
    indexed = literature.set_index("key", drop=False)

    def cite(*keys: str) -> str:
        return _parenthetical_citations(indexed.loc[list(keys)])

    return cite


def _validate_literature(path: Path, literature: pd.DataFrame) -> None:
    text = _document_text(path)
    body = text.split("\nLiterature cited\n", 1)[0]
    labels = set(literature.citation)
    uncited = sorted(label for label in labels if label not in body)
    if uncited:
        raise ValueError(f"{path.name}: references not cited in text: {uncited}")
    groups = re.findall(r"\(([^()]*\d{4})\)", body)
    in_text = {
        item.strip()
        for group in groups
        for item in group.split(";")
        if re.fullmatch(r"[A-Z][^;]* \d{4}", item.strip())
    }
    orphaned = sorted(in_text - labels)
    if orphaned:
        raise ValueError(f"{path.name}: in-text citations without references: {orphaned}")
    for group in groups:
        items = [item.strip() for item in group.split(";")]
        if all(item in labels for item in items):
            years = [int(item[-4:]) for item in items]
            if years != sorted(years):
                raise ValueError(f"{path.name}: citations not chronological: {group}")


def _main_manuscript(
    root: Path,
    summary: pd.DataFrame,
    information: pd.DataFrame,
    coverage: pd.DataFrame,
    literature: pd.DataFrame,
    checkpoint: dict,
) -> Path:
    output = root / "manuscript/INLINE_MANUSCRIPT_HUMAN_REVIEW_DRAFT.docx"
    utility_weights = json.loads((root / "config/analysis.json").read_text())[
        "information_utility_weights"
    ]
    document = Document()
    title = TITLE
    _configure_document(document, title)
    _add_title(document, title, "Conservation Biology Contributed Paper — anonymous draft")
    _add_review_notice(document)

    reversals, advantages, indeterminate = _sweep_counts(root)
    audit_replications = int(
        pd.read_csv(root / "results/convergence_audit.csv").n.max()
    )
    information_effect = information[
        information.metric == "decision_utility"
    ].iloc[0]
    abstract = (
        "Conservation decisions increasingly link habitat protection, managed care, "
        "movement, breeding, survey, and biobanking, yet the conditions under which "
        "dynamic networks outperform fixed strategies remain unclear. We developed "
        "an individual-based, annual simulation of eight synthetic species across "
        "nine wild and managed node types for 100 years. Five species were initially "
        "known and three were latent to feasible decision makers. Eight policies "
        "were compared under finite budgets, environmental change, noisy delayed "
        "observations, welfare costs, genetics, disease, human engagement, and "
        "origin fidelity. The production design used "
        f"{checkpoint['selected_replications']} paired-seed replications per policy, "
        "Latin-hypercube sensitivity analysis, five two-factor sweeps, feature "
        "ablations, adverse conditions, and a current-state-information comparator. "
        "Habitat-first policy S6 had mean year-100 survival "
        f"{_interval(summary, 'S6', 'survival_100')}, "
        f"whereas the hybrid policy S7 had {_interval(summary, 'S7', 'survival_100')}. "
        f"Robust sweep-cell contrasts favored S7 in {advantages} cells and S6 in "
        f"{reversals}; {indeterminate} cells were directionally indeterminate. "
        "S7 produced mean individual welfare "
        f"{_interval(summary, 'S7', 'mean_individual_welfare')}; lifetime captivity "
        f"S2 produced {_interval(summary, 'S2', 'mean_individual_welfare')}. "
        "The paired current-state-information effect on the prespecified utility "
        f"scale was {information_effect.estimate:.3f} "
        f"[{information_effect.lower:.3f}, {information_effect.upper:.3f}]. "
        "Dynamic conservation therefore had no universal advantage: benefits "
        "depended on habitat change, survey cost, movement harm, managed-care "
        "benefit, and the objectives being valued. The framework identifies "
        "assumption-dependent reversal boundaries rather than prescribing one "
        "institutional form."
    )
    abstract_words = len(re.findall(r"\b[\w–-]+\b", abstract))
    if abstract_words > 300:
        raise ValueError(f"Abstract is {abstract_words} words")
    _add_heading(document, "Abstract")
    _add_paragraph(document, abstract)
    _add_paragraph(
        document,
        "Keywords: biodiversity discovery; conservation decision analysis; "
        "ex situ conservation; partial observability; policy reversal; reintroduction; "
        "value of information",
        bold_prefix="Keywords:",
    )
    _add_paragraph(
        document,
        "Article impact statement: Hidden biodiversity and shifting habitat can "
        "reverse which mix of protection, managed care, and movement performs best.",
        bold_prefix="Article impact statement:",
    )

    _add_heading(document, "Introduction")
    cite = _citer(literature)
    _add_paragraph(
        document,
        "Conservation portfolios operate across environments rather than within a "
        "single protected place. About one in seven threatened terrestrial "
        f"vertebrate species is held in captivity {cite('conde_2011_science')}, and "
        "proposals to integrate captive and wild management have argued for a larger "
        "ex situ role alongside dominant in situ approaches "
        f"{cite('pritchard_2012', 'redford_2012')}. Threatened species are unevenly "
        "represented in zoo collections, and managing viable zoo metapopulations "
        f"requires coordination among many institutions {cite('conde_2013_plos')}. "
        "Zoos and aquariums can supply source animals for reintroduction "
        f"{cite('gilbert_2017')} within a spectrum of conservation translocations "
        "that extends from reinforcement and reintroduction to introductions "
        f"outside the indigenous range {cite('seddon_2014')}. Habitat protection, "
        "zoological and breeding facilities, rehabilitation, assisted movement, "
        "prerelease environments, survey, and biobanking can therefore alter "
        "different components of persistence."
    )
    _add_paragraph(
        document,
        "Each component also carries documented costs. Captive breeding has faced "
        "poor reintroduction success, high costs, domestication, disease outbreaks, "
        "and preemption of other recovery techniques "
        f"{cite('snyder_1996')}. Genetic adaptation to captivity is generally "
        f"deleterious when populations are returned to the wild {cite('frankham_2008')}; "
        "in steelhead trout, captive rearing reduced subsequent reproductive success "
        "in the wild by approximately 40% per captive-reared generation "
        f"{cite('araki_2007')}. Potential welfare problems were recorded in 67% of "
        "199 reviewed reintroduction projects, whereas welfare was explicitly "
        f"addressed in 6% {cite('harrington_2013')}, and harms to individual welfare "
        f"can also hamper conservation goals {cite('beausoleil_2018')}. Under climate "
        "change, moving species outside historic ranges has been proposed to reduce "
        f"biodiversity loss {cite('hoegh_guldberg_2008')}, but assisted migration has "
        f"required an explicit framework for debate {cite('mclachlan_2007')}."
    )
    _add_paragraph(
        document,
        "These effects occur under incomplete knowledge: managers may observe "
        "abundance or habitat with error and delay, while some taxa remain "
        "undetected, undescribed, or unrecognized as threatened. One estimate "
        "suggests that about 86% of existing species await description "
        f"{cite('mora_2011')}, although richness estimates and the risk of extinction "
        f"before naming are contested {cite('costello_2013')}. Undescribed species "
        "may account for 15–59% of recent extinctions, depending on taxon and region "
        f"{cite('tedesco_2014')}; such losses have been termed dark extinction "
        f"{cite('boehm_cronk_2021')}, and the proportion threatened among newly "
        f"described vertebrates has increased over time {cite('liu_2022')}. A "
        "strategy that protects known populations efficiently can therefore differ "
        "from one that preserves options for unknown biodiversity."
    )
    _add_paragraph(
        document,
        "Complete-text case evidence illustrates complementary requirements rather "
        "than a universal recipe. A Przewalski's horse project progressed from "
        "captive-born release toward standardized monitoring and broader ecosystem "
        "work. Activity-specific niche models evaluated with withheld data identified "
        "candidate California condor reintroduction areas, while the authors retained "
        "the need for ground reconnaissance and assessment of unmodeled threats. In a "
        "reintroduced southern California condor population, anthropogenic-material "
        f"ingestion contributed to nestling mortality and low nest success "
        f"{cite('kaczensky_2007', 'mee_2007', 'delia_2015')}. Ecological decision "
        "theory under partial observability separates hidden states from observations "
        "and tracks belief states, but adds substantial computational and "
        f"interpretive complexity {cite('williams_brown_2022')}. Such formulations "
        "have shown that managing for a cryptic threatened species can be optimal "
        f"even when its presence is uncertain {cite('chades_2008')} and formalize "
        "trade-offs between management and surveillance under imperfect observation "
        f"{cite('chades_2021')}."
    )
    _add_paragraph(
        document,
        "Existing debates can be distorted by treating survival, wild restoration, "
        "experienced welfare, genetic representation, disease, cost, and future "
        "option value as interchangeable. A dynamic network may improve one outcome "
        "while worsening another, and movement may be beneficial when future habitat "
        "shifts but harmful when origin habitat remains excellent or transfer risk "
        "is high. The relevant scientific problem is consequently conditional: when "
        "do dynamic networks change outcomes relative to fixed strategies, and when "
        "do those differences disappear or reverse?"
    )
    _add_paragraph(
        document,
        "We addressed this question with a transparent synthetic decision experiment. "
        "The contribution is not a calibrated population forecast or a universal "
        "recommendation. It is an integrated test bed in which latent biodiversity, "
        "individual heterogeneity, partial observability, wild and managed nodes, "
        "finite resources, and multiple outcomes can be varied together. We expected "
        "policy rankings to reverse across environmental and institutional conditions "
        "and treated those reversals, rather than a base-case winner, as the principal "
        "result."
    )

    _add_heading(document, "Methods")
    _add_heading(document, "Simulation scope and state", level=2)
    _add_paragraph(
        document,
        "The model followed eight synthetic species for 100 annual steps. Five "
        "species were initially represented in the observation state; three existed "
        "in the true ecological state but were latent to feasible decision makers. "
        "Latent species progressed through existence, detection, description, "
        "monitoring, threat recognition, and protection. A latent population could "
        "become extinct before description, producing the extinction-before-discovery "
        "endpoint. "
        f"{TABLE['parameters']} lists the base parameters used in the production "
        "run. State definitions, observation rules, budget constraints, "
        "environmental transitions, disease, welfare, genetic proxies, biobanking, "
        "and policy engines are specified in the reproducibility repository."
    )
    _main_table(document, root, "parameters")
    _add_paragraph(
        document,
        "Individuals carried sex, age or life stage, origin, current node, founder or "
        "lineage, inbreeding proxy, health, infection, wildness, habituation, stress "
        "tolerance, foraging ability, release readiness, and alive status. Species "
        "traits governed growth, mortality, niche breadth, habitat specificity, "
        "climate and anthropogenic sensitivity, dispersal, detectability, fertility, "
        "social dependence, migration need, and captivity tolerance. Population "
        "transitions combined survival, reproduction, movement, infection, welfare, "
        "and environmental change while enforcing a maximum living population."
    )
    _add_paragraph(
        document,
        "These processes were represented qualitatively with synthetic values. "
        "Relocation stress reflects the view that stress is an inevitable component "
        f"of animal translocation {cite('dickens_2010')}; release readiness reflects "
        "analyses of captive experience and reintroduction survival "
        f"{cite('jule_2008')}; and transfer-associated infection reflects the disease "
        f"risks of wildlife translocation {cite('kock_2010')}. Founder and "
        "inbreeding proxies were retained because genetic-diversity targets and "
        "indicators were undeveloped in earlier global biodiversity strategies "
        f"{cite('hoban_2020')} and because minimizing genetic adaptation is a "
        f"recognized aim of captive breeding {cite('williams_hoffman_2009')}. The "
        "welfare index is synthetic and is not an implementation of multidomain "
        f"assessment such as the Five Domains Model {cite('mellor_2020')}. Biobank "
        "value is stylized; cryobanking combined with assisted reproduction can "
        f"reinstate lost genetic diversity {cite('bolton_2022')}. The engagement "
        "feedback is also synthetic: increased biodiversity understanding over zoo "
        f"and aquarium visits {cite('moss_2015')} does not establish the "
        "habitat-protection effect assumed in the model."
    )
    _add_paragraph(
        document,
        "Nine node types represented original wild habitat, alternative wild habitat, "
        "protected reserve, conventional zoo, conservation zoo, breeding center, "
        "semi-wild or prerelease facility, rehabilitation facility, and biobank. "
        f"{FIGURE['network']} shows the modeled network and potential movement paths."
    )
    _main_figure(document, root, "network")

    _add_heading(document, "Observation and policy decisions", level=2)
    _add_paragraph(
        document,
        "Feasible policies received noisy, delayed observations rather than true "
        "habitat quality or latent-species state. The model distinguished the true "
        "number of species into observed and latent components. Survey spending "
        "altered discovery progression; protection could begin only after sufficient "
        "recognition. An oracle analysis exposed true current abundance and habitat "
        "quality to the same S7 action class. It did not expose future environmental "
        "draws, solve a partially observable Markov decision process (POMDP), or "
        "select actions using realized future utility. The contrast was motivated by "
        "value-of-information reasoning, which asks whether reducing uncertainty "
        f"would change a management decision {cite('canessa_2015')}, but it is not a "
        "value-of-information calculation."
    )
    _add_paragraph(
        document,
        "Eight policies ranged from no intervention through fixed rules to dynamic "
        f"allocation ({TABLE['policies']}). S0 applied no intervention. S1 phased down ex situ "
        "management, S2 emphasized lifetime captivity, S3 returned individuals to "
        "origin when feasible, S4 used state-dependent circulation, S5 used greedy "
        "expected utility without automatic origin privilege, S6 emphasized habitat, "
        "and S7 combined habitat, survey, selective circulation, breeding, genetics, "
        "biobanking, and engagement with a forward-quality proxy. Annual spending "
        "could not exceed available resources, reflecting the importance of "
        "management cost and probability of success in allocation among threatened "
        f"species {cite('joseph_2009')}."
    )
    _main_table(document, root, "policies")

    _add_heading(document, "Outcomes and analysis", level=2)
    _add_paragraph(
        document,
        "Outcomes were retained separately: survival and species retention, wild "
        "abundance, extinction before discovery, cumulative welfare-person-years, "
        "mean individual welfare, effective-founder and inbreeding proxies, disease, "
        "successful reintroduction, relocation burden, habitat integrity, future "
        "option value, engagement, expenditure, and cost per species retained. The "
        "prespecified information-value utility weighted year-100 survival "
        f"({utility_weights['survival_100']:.2f}), future option value "
        f"({utility_weights['future_option_value']:.2f}), habitat integrity "
        f"({utility_weights['habitat_integrity']:.2f}), avoidance of extinction "
        f"before discovery ({utility_weights['ebd_avoidance']:.2f}), and mean "
        "individual welfare "
        f"({utility_weights['mean_individual_welfare']:.2f}). This utility was "
        "used only for the oracle contrast and not to declare an overall policy winner. "
        "Policies were instead compared by Pareto nondominance, which displays "
        "trade-offs among conflicting objectives without weighting them into a "
        f"single value {cite('kennedy_2008')}."
    )
    _add_paragraph(
        document,
        f"The production analysis used {checkpoint['selected_replications']} "
        "replications per policy with shared seed labels. Because policy actions "
        "change later random draws, these are paired seeds rather than mathematically "
        "exact common random streams. Intervals are percentiles of the simulation "
        "replicate distribution, not empirical confidence intervals. Sensitivity "
        "analysis used a Latin-hypercube design over 12 parameters. Five two-factor "
        "sweeps varied habitat deterioration by relocation stress, latent fraction by "
        "survey cost, climate velocity by assisted-colonization risk, captivity burden "
        "by breeding benefit, and engagement effect by habitat-protection efficiency. "
        "Nine feature ablations and six adverse-condition designs evaluated whether "
        "differences disappeared or reversed. The common replication count was fixed "
        f"after a nested convergence audit through {audit_replications} paired "
        "replications, and seeds were recorded with each policy outcome. Stochastic "
        "replicates were averaged within each Latin-hypercube design before Spearman "
        f"association ({TABLE['sensitivity']}) and partial-rank interaction "
        f"screening ({TABLE['interactions']}). Cell-level two-factor sign stability "
        "was evaluated with nested replicate prefixes."
    )
    _main_table(document, root, "sensitivity")
    _main_table(document, root, "interactions")
    _add_paragraph(
        document,
        "The empirical component was deliberately separate from the synthetic model. "
        "Persisted Global Biodiversity Information Facility (GBIF) API responses "
        "supplied taxonomy matches and occurrence-year "
        f"coverage aggregates for three conservation taxa ({TABLE['coverage']}); "
        "they did not calibrate "
        "abundance, occupancy, demographic rates, or intervention effects. Persisted "
        "Bibliographic, repository, and publisher snapshots supported an evidence-level "
        f"audit ({TABLE['evidence']}). Complete text was inspected for the four case "
        "and decision-theory articles. Other background citations were verified "
        "against exact-DOI Crossref metadata, and statements attributed to them were "
        "restricted to their retained author abstracts or article titles. Five "
        "formerly selected sources without verified text were excluded from the "
        "citation set. The literature did not calibrate the model."
    )
    _main_table(document, root, "coverage")
    _main_table(document, root, "evidence")

    _add_heading(document, "Results")
    _add_heading(document, "Base-case policy outcomes", level=2)
    _add_paragraph(
        document,
        f"No single policy dominated every outcome ({TABLE['outcomes']}). S6 had the largest mean "
        f"year-100 survival, {_interval(summary, 'S6', 'survival_100')}, and retained "
        f"{_interval(summary, 'S6', 'species_retained_100')} species. S7 had survival "
        f"{_interval(summary, 'S7', 'survival_100')} and retained "
        f"{_interval(summary, 'S7', 'species_retained_100')} species. S4 had survival "
        f"{_interval(summary, 'S4', 'survival_100')}. S0 reached zero mean year-100 "
        "survival in the directional-change base scenario, but that outcome is a "
        "property of the configured experiment rather than an empirical estimate. "
        f"Multiobjective Pareto status is shown in {FIGURE['pareto']}."
    )
    _main_table(document, root, "outcomes")
    _main_figure(document, root, "pareto")
    _add_paragraph(
        document,
        "Temporal trajectories showed divergence among policies as directional "
        f"environmental change accumulated ({FIGURE['trajectories']}). Policies "
        "with habitat investment maintained higher average survival than no "
        "intervention in the base scenario, "
        "whereas fixed return to origin did not preserve species through year 100. "
        "Managed-care and dynamic policies differed in how persistence was distributed "
        "between wild and managed nodes."
    )
    _main_figure(document, root, "trajectories")

    _add_heading(document, "Welfare, hidden biodiversity, and information", level=2)
    _add_paragraph(
        document,
        f"S7 mean individual welfare was {_interval(summary, 'S7', 'mean_individual_welfare')}; "
        "the corresponding estimate for lifetime captivity S2 was "
        f"{_interval(summary, 'S2', 'mean_individual_welfare')}. S7's extinction-before-"
        f"discovery fraction was {_interval(summary, 'S7', 'fraction_ebd')}, compared "
        f"with {_interval(summary, 'S5', 'fraction_ebd')} for S5 and "
        f"{_interval(summary, 'S6', 'fraction_ebd')} for S6; extinction before "
        "discovery across latent-fraction and survey-cost conditions is shown in "
        f"{FIGURE['ebd']}. These outcomes remained "
        "distinct because high persistence in managed care did not necessarily imply "
        "restoration of free-living populations, high individual welfare, or low "
        f"extinction before discovery ({FIGURE['tradeoffs']}). Normalized "
        "cross-objective values are descriptive and do not imply cardinal "
        "comparability among objectives."
    )
    _main_figure(document, root, "ebd")
    _main_figure(document, root, "tradeoffs")
    _add_paragraph(
        document,
        "Supplying true current state to the same S7 action class changed the "
        "prespecified decision utility by "
        f"{information_effect.estimate:.3f} "
        f"[{information_effect.lower:.3f}, {information_effect.upper:.3f}] "
        "on average. The "
        "paired effect could be positive or negative because the decision rule was "
        "not reoptimized and no hindsight action selection was applied. This estimate "
        "is conditional on the policy class and normative utility weights; it is not "
        "theoretical EVPI or an intrinsic monetary or ecological value of biodiversity "
        "information. The separation of true state, observations, actions, and "
        f"outcomes is shown in {FIGURE['observation']}."
    )
    _main_figure(document, root, "observation")

    _add_heading(document, "Reversal and robustness conditions", level=2)
    _add_paragraph(
        document,
        f"Across the prespecified two-factor cells, robust survival contrasts favored "
        f"S7 in {advantages} cells and S6 in {reversals}; {indeterminate} cells "
        "were indeterminate under the registered sign-stability and Monte Carlo "
        "error criterion. "
        f"Selected reversal surfaces are shown in {FIGURE['reversal']}. Habitat deterioration, "
        "relocation stress, survey cost, assisted-colonization risk, captivity burden, "
        "breeding benefit, engagement, and protection efficiency jointly changed the "
        "direction or magnitude of policy contrasts. Survey-cost "
        f"exploration–exploitation patterns are shown in {FIGURE['exploration']}."
    )
    _main_figure(document, root, "reversal")
    _main_figure(document, root, "exploration")
    _add_paragraph(
        document,
        "Adverse-condition designs clarified mechanisms. Stable excellent origin "
        "habitat and high relocation harm weakened unnecessary movement; severe "
        "captivity burden reduced welfare under managed care; zero breeding advantage "
        "removed a major pathway through which ex situ investment could aid "
        "persistence; and efficient restoration strengthened habitat-centered "
        "strategies. Removing latent species eliminated extinction-before-discovery "
        "by construction. Removing origin fidelity changed dynamic allocation without "
        "making any one dynamic policy uniformly preferable. Feature ablations "
        "removed human feedback, latent species, genetics, climate, biobanking, "
        "disease, welfare, partial observability, and origin fidelity "
        f"({TABLE['robustness']})."
    )
    _main_table(document, root, "robustness")

    _add_heading(document, "Empirical source audit", level=2)
    coverage_text = "; ".join(
        f"{row.species}, {int(row.gbif_query_total):,} coordinate-bearing present "
        f"records ({int(row.records_2015_2025):,} dated 2015–2025)"
        for row in coverage.itertuples()
    )
    _add_paragraph(
        document,
        f"The persisted GBIF aggregate queries returned {coverage_text}. Differences "
        "can reflect observation effort, accessibility, digitization, taxonomy, and "
        "data publishing. They cannot be interpreted as relative abundance, occupancy, "
        "population trend, or conservation effectiveness. The four case and "
        "decision-theory publications were verified by DOI, title, year, journal, and "
        "authorship, and their complete methods, results, and discussion text was "
        "appraised for the bounded claims in the Introduction."
    )

    _add_heading(document, "Discussion")
    _add_paragraph(
        document,
        "The central result was conditionality rather than dominance. Habitat-first "
        "S6 had the highest mean survival in the base configuration, while S7 had "
        "higher mean welfare and future option value, but their survival ordering "
        "reversed repeatedly across prespecified conditions. Dynamic conservation "
        "networks therefore cannot be evaluated by movement or institutional "
        "connectivity alone. Their effect depends on whether changing habitat creates "
        "a real future-location benefit large enough to exceed relocation mortality, "
        "stress, pathogen transmission, maladaptation, captivity burden, survey cost, "
        "and the opportunity cost of diverting resources from habitat. This "
        "conditionality is consistent with arguments that managed relocation involves "
        "interacting, value-laden considerations and decisions under imperfect "
        f"information {cite('richardson_2009')} and that climate-driven translocation "
        "challenges the aim of recreating past ecological communities "
        f"{cite('thomas_2011')}, which bears directly on origin fidelity. The deterministic "
        "origin-to-future habitat-quality shift under zero intervention and zero "
        f"stochastic shock is shown in {FIGURE['origin_future']}."
    )
    _main_figure(document, root, "origin_future")
    _add_paragraph(
        document,
        "The framework's scientific contribution is to place wild habitat, managed "
        "care, transitional facilities, rehabilitation, survey, and biobanking in one "
        "partially observed allocation problem while keeping welfare, genetics, "
        "disease, information, and cost as explicit outcomes. Latent taxa are modeled "
        "as real populations hidden from decisions rather than as missing labels, so "
        "discovery delay can interact with extinction. The survey–protection balance "
        "is a form of the trade-off between management performance and learning "
        f"addressed by active adaptive management {cite('mccarthy_possingham_2007')}. "
        "The current-state-information "
        "comparator further shows that information effects arise only through a "
        "decision rule and objective function; its sign and magnitude under alternative "
        "utility weights (equal-component, persistence-, welfare-, future-option-, "
        "and extinction-before-discovery-priority weights) are reported in "
        f"{TABLE['information_weights']}. These diagnostics change only the declared "
        "aggregation weights, not simulated actions or outcomes."
    )
    _main_table(document, root, "information_weights")
    _add_paragraph(
        document,
        "This structure separates persistence from restoration. A captive population "
        "can preserve living individuals or founder representation without restoring "
        "a free-living population, whereas rapid release can increase wild abundance "
        "while imposing movement mortality or disease. Likewise, cumulative "
        "welfare-person-years rewards both survival duration and experienced "
        "conditions, while mean individual welfare isolates average conditions among "
        "living individuals. Reporting these outcomes separately prevents a single "
        "weighted score from concealing normative trade-offs."
    )
    _add_paragraph(
        document,
        "The results do not establish that any real zoo, reserve, translocation "
        "program, breeding center, or biobank should expand or contract. Species are "
        "synthetic, demographic and genetic processes are simplified, policy shares "
        "are fixed, and the S7 forward-quality proxy is not a solved partially "
        "observable Markov decision process. Disease, catastrophes, social behavior, "
        "and biobank viability are stylized, although modeled integration of "
        "biobanked sperm into amphibian captive breeding reduced programme cost and "
        f"made genetic-retention targets feasible {cite('howell_2021')}. Replicate "
        "intervals characterize model "
        "variation, not sampling uncertainty in a real population. Paired seed labels "
        "also do not ensure identical random streams after policies change system state."
    )
    _add_paragraph(
        document,
        "Applied use would require preregistered taxon-specific calibration, pedigree "
        "or genomic dynamics, epidemiological transmission models, facility and "
        "landscape constraints, stakeholder-derived multiattribute utilities, and "
        "full-text evidence synthesis. Observation systems should also be evaluated "
        "for bias: the GBIF audit demonstrates only data availability, and high record "
        "counts may primarily reflect survey and publishing effort."
    )
    _add_paragraph(
        document,
        "Within those limits, the model supports a practical research principle. "
        "Conservation network design should be tested against explicit reversal "
        "conditions before base-case rankings are generalized. Movement is most "
        "defensible when future habitat gains and managed-care benefits exceed welfare, "
        "disease, and opportunity costs; habitat-centered strategies become more "
        "attractive when origin quality is stable, restoration is efficient, or "
        "movement is harmful, so evidence on the performance and potential of "
        f"protected areas {cite('watson_2014')} bears directly on the habitat-first "
        "comparator. Survey and information investment matter most when they "
        "can alter a feasible action before hidden populations disappear."
    )

    _add_heading(document, "Acknowledgments")
    _add_paragraph(
        document,
        "[HUMAN AUTHORS: add only verified acknowledgments after obtaining any "
        "required permission. Keep identifying information out of the anonymous file.]"
    )
    _add_heading(document, "Data and code availability")
    _add_paragraph(
        document,
        "Synthetic inputs, code, seed assignments, result tables, and figure scripts "
        "are organized in the reproducibility repository. Public-source API responses "
        "are retained as immutable timestamped snapshots with URLs, retrieval "
        "conditions, sizes, SHA-256 checksums, license notes, and completeness "
        "statements. Before submission, human authors must add a durable repository "
        "citation, version or DOI, license, and anonymous review link. The repository "
        "contains GBIF aggregate year-facet responses and bibliographic metadata, not "
        "individual occurrence records. Publisher full text retained for appraisal is "
        "not redistributed in the public or anonymous packages."
    )
    _add_heading(document, "Artificial-intelligence use disclosure")
    _add_paragraph(
        document,
        "Devin, an artificial-intelligence software engineering system developed by "
        "Cognition AI, was used to implement and test code, organize public-source "
        "metadata, execute analyses, generate code-derived displays, and prepare this "
        "human-review draft. The AI did not create or alter empirical source data and "
        "is not an author. Human authors must independently inspect and revise all "
        "prose, citations, analyses, assumptions, figures, and tables; determine the "
        "scientific interpretation; and accept responsibility for the submitted work. "
        "[HUMAN AUTHORS: verify the exact product/model description and final journal "
        "placement before submission.]"
    )
    _add_heading(document, "Literature cited")
    _add_paragraph(
        document,
        "Complete article text was appraised for the four case and decision-theory "
        "entries; the remaining entries were verified against exact-DOI Crossref "
        "metadata and are cited only for statements supported by their author "
        "abstracts or titles. Final claim wording and citation approval remain the "
        "responsibility of the human authors."
    )
    ordered = literature.assign(
        sort_key=literature.reference_authors.str.casefold()
    ).sort_values(["sort_key", "year"])
    for reference in _reference_text(ordered):
        _add_paragraph(document, reference)

    document.save(output)
    return output


def _editable_tables(root: Path) -> Path:
    output = root / "manuscript/EDITABLE_TABLES.docx"
    document = Document()
    _configure_document(document, "Editable tables")
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11)
    section.page_height = Inches(8.5)
    _add_title(document, "Editable tables", "Anonymous submission-support file")
    for number, filename, caption in MAIN_TABLES:
        _add_table(
            document,
            pd.read_csv(root / "tables" / filename),
            number,
            caption,
            font_size=7 if number in COMPACT_TABLES else 8,
        )
    document.save(output)
    return output


def _png_dimensions(path: Path) -> tuple[int, int]:
    with path.open("rb") as stream:
        header = stream.read(24)
    if header[:8] != b"\x89PNG\r\n\x1a\n":
        raise ValueError(f"Not a PNG: {path}")
    return struct.unpack(">II", header[16:24])


def _figure_deck(root: Path) -> Path:
    legacy_output = root / "manuscript/EDITABLE_FIGURES.pptx"
    legacy_output.unlink(missing_ok=True)
    output = root / "manuscript/FIGURE_DECK.pptx"
    presentation = Presentation()
    presentation.slide_width = PptxInches(13.333)
    presentation.slide_height = PptxInches(7.5)
    presentation.core_properties.title = "Submission figure deck"
    presentation.core_properties.subject = (
        "Editable slide titles and captions with image-based figure artwork"
    )
    presentation.core_properties.author = ""
    presentation.core_properties.last_modified_by = ""
    presentation.core_properties.comments = ""
    blank_layout = presentation.slide_layouts[6]
    for number, filename, caption in MAIN_FIGURES:
        slide = presentation.slides.add_slide(blank_layout)
        title_box = slide.shapes.add_textbox(
            PptxInches(0.45), PptxInches(0.15), PptxInches(12.43), PptxInches(0.55)
        )
        title_frame = title_box.text_frame
        title_frame.text = f"{number}. {caption.split('.')[0]}."
        title_paragraph = title_frame.paragraphs[0]
        title_paragraph.alignment = PP_ALIGN.CENTER
        title_paragraph.runs[0].font.name = "Arial"
        title_paragraph.runs[0].font.size = PptxPt(24)
        title_paragraph.runs[0].font.bold = True
        image = root / "figures" / filename
        width_px, height_px = _png_dimensions(image)
        max_width = 12.0
        max_height = 5.55
        ratio = min(max_width / width_px, max_height / height_px)
        image_width = width_px * ratio
        image_height = height_px * ratio
        image_left = (13.333 - image_width) / 2
        image_top = 0.75 + (max_height - image_height) / 2
        slide.shapes.add_picture(
            str(image),
            PptxInches(image_left),
            PptxInches(image_top),
            width=PptxInches(image_width),
            height=PptxInches(image_height),
        )
        caption_box = slide.shapes.add_textbox(
            PptxInches(0.55), PptxInches(6.42), PptxInches(12.23), PptxInches(0.8)
        )
        caption_frame = caption_box.text_frame
        caption_frame.word_wrap = True
        caption_frame.text = f"{number}. {caption}"
        caption_paragraph = caption_frame.paragraphs[0]
        caption_paragraph.alignment = PP_ALIGN.LEFT
        caption_paragraph.runs[0].font.name = "Arial"
        caption_paragraph.runs[0].font.size = PptxPt(12)
        note_box = slide.shapes.add_textbox(
            PptxInches(0.55), PptxInches(7.22), PptxInches(12.23), PptxInches(0.18)
        )
        note_frame = note_box.text_frame
        note_frame.text = (
            "Image-based artwork; edit titles/captions here and regenerate artwork "
            "with make figures."
        )
        note_paragraph = note_frame.paragraphs[0]
        note_paragraph.alignment = PP_ALIGN.RIGHT
        note_paragraph.runs[0].font.name = "Arial"
        note_paragraph.runs[0].font.size = PptxPt(8)
    presentation.save(output)
    return output


def _cover_page(root: Path) -> Path:
    output = root / "manuscript/COVER_PAGE_TEMPLATE.docx"
    document = Document()
    _configure_document(document, "Cover page template")
    document.core_properties.subject = "Identifying cover page template"
    _add_title(document, "Cover page template")
    _add_paragraph(
        document,
        "Replace every bracketed placeholder. Do not upload this file as the anonymous "
        "manuscript.",
    )
    fields = [
        ("Article title", "[final approved title]"),
        ("Article type", "Contributed Paper"),
        ("Running head", "Dynamic conservation networks"),
        ("Word count", "[Abstract through Acknowledgments]"),
        ("Authors", "[full publication names in approved order]"),
        ("Affiliations", "[verified affiliations]"),
        ("Corresponding author", "[name, address, email, and required contact details]"),
        ("ORCID identifiers", "[verified identifiers, if applicable]"),
        ("Author contributions", "[verified CRediT roles]"),
        ("Funding", "[verified funder and award identifiers, or explicit none]"),
        ("Competing interests", "[verified declaration]"),
        ("Acknowledgments", "[verified names and permissions]"),
        ("Data and code availability", "[final archival citation and anonymous link]"),
        ("AI-use disclosure", "[authors' reviewed and approved final statement]"),
    ]
    for label, value in fields:
        _add_paragraph(document, f"{label}: {value}", bold_prefix=f"{label}:")
    document.save(output)
    return output


def _cover_letter_text(root: Path) -> str:
    values = pd.read_csv(root / "results/manuscript_values.csv").set_index("value_id")
    pareto = pd.read_csv(root / "results/pareto_front.csv")
    words = ["zero", "one", "two", "three", "four", "five", "six", "seven", "eight"]
    s6 = values.loc["s6_survival_100", "estimate"]
    s7 = values.loc["s7_survival_100", "estimate"]
    efficient = words[int(pareto.pareto_efficient.sum())]
    total = words[len(pareto)]
    paragraphs = [
        "[Date]",
        "Editors\nConservation Biology",
        "Dear Editors,",
        f"We submit the manuscript \u201c{TITLE}\u201d for consideration as a "
        "Contributed Paper in Conservation Biology.",
        "When should an animal be held in managed care, moved, returned to its "
        "origin, or left to habitat protection, and how does the answer change "
        "when some of the species at stake have not yet been detected? Conservation "
        "practice increasingly links reserves, zoos, breeding centers, prerelease "
        "facilities, rehabilitation, survey, and biobanking, yet these components "
        "are usually debated one institution or one objective at a time. Our study "
        "asks a conditional question instead: under which ecological, informational, "
        "and welfare conditions does a dynamic conservation network outperform a "
        "fixed strategy, and when does that advantage disappear or reverse?",
        "We address this with a transparent individual-based simulation spanning "
        "100 years, eight synthetic species of which three are initially "
        "undetected, and nine wild and managed node types. Policies are compared "
        "under finite budgets, directional environmental change, noisy and delayed "
        "observation, welfare, genetics, disease, human engagement, and origin "
        "fidelity, and each objective is kept separate rather than collapsed into "
        "a single score.",
        "The central result is that no conservation network is universally best. "
        "A habitat-first policy retained slightly higher mean year-100 survival "
        f"than a hybrid dynamic policy ({s6:.3f} vs. {s7:.3f}), yet {efficient} of "
        f"{total} policies were Pareto-nondominated once survival, welfare, "
        "genetic outcome, extinction before discovery, cost, and future option "
        "value were evaluated jointly. Which policy is favored reverses across "
        "identifiable boundaries in habitat change, survey cost, movement harm, "
        "and managed-care benefit. These boundaries, rather than a ranking, are "
        "the contribution.",
        "We believe the work suits Conservation Biology because its relevance "
        "extends beyond any single species, site, or institution. It offers a "
        "reusable way to state the conditions under which integrated in situ and "
        "ex situ management is worth its costs, including costs borne by "
        "individual animals and by species not yet known to science. The results "
        "are synthetic conditional experiments, not empirical forecasts, and the "
        "manuscript makes no general recommendation for or against zoos, "
        "translocation, return to origin, or habitat-first conservation.",
        "All code, configuration, and generated outputs needed to reproduce every "
        "number, figure, and table are available in a public repository, cited in "
        "the data availability statement. The manuscript has been prepared for "
        "double-blind review. It is not under consideration elsewhere, and all "
        "authors have approved its submission.",
        "Sincerely,",
        "[Corresponding author name]\n[Affiliation]\n[Email address]",
    ]
    return "\n\n".join(paragraphs) + "\n"


def _cover_letter(root: Path) -> Path:
    text = _cover_letter_text(root)
    (root / "manuscript/COVER_LETTER.txt").write_text(text, encoding="utf-8")
    output = root / "manuscript/COVER_LETTER.docx"
    document = Document()
    _configure_document(document, "Cover letter")
    document.core_properties.subject = "Cover letter to the editors"
    normal = document.styles["Normal"]
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    normal.paragraph_format.space_after = Pt(10)
    normal.paragraph_format.widow_control = True
    for block in text.strip().split("\n\n"):
        paragraph = document.add_paragraph()
        lines = block.split("\n")
        for index, line in enumerate(lines):
            run = paragraph.add_run(line)
            if index < len(lines) - 1:
                run.add_break()
    document.save(output)
    return output


def _document_text(path: Path) -> str:
    document = Document(path)
    lines = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(lines)


def _validate_references(path: Path, figures: list[str], tables: list[str]) -> None:
    text = _document_text(path)
    for labels in [figures, tables]:
        first_positions = []
        for label in labels:
            positions = [match.start() for match in re.finditer(re.escape(label), text)]
            if len(positions) < 2:
                raise ValueError(f"{path.name}: {label} lacks citation or caption")
            first_positions.append(positions[0])
        if first_positions != sorted(first_positions):
            raise ValueError(f"{path.name}: references are not sequential: {labels}")


def _validate_office_file(path: Path) -> None:
    with zipfile.ZipFile(path) as archive:
        invalid = archive.testzip()
        if invalid:
            raise ValueError(f"{path.name}: corrupt member {invalid}")


def _scan_anonymous_office_file(path: Path) -> None:
    forbidden = ["bougtoir", "tatsuki", "onishi", "@gmail", "/home/", "app.devin.ai"]
    with zipfile.ZipFile(path) as archive:
        text = "\n".join(
            archive.read(name).decode("utf-8", errors="ignore").lower()
            for name in archive.namelist()
            if name.endswith((".xml", ".rels"))
        )
    findings = [term for term in forbidden if term in text or term in path.name.lower()]
    if findings:
        raise ValueError(f"{path.name}: identifying content found: {findings}")


def _count_words(text: str) -> int:
    return len(re.findall(r"\b[\w–-]+\b", text))


def _manuscript_word_count(path: Path) -> int:
    document = Document(path)
    started = False
    acknowledgments_heading_level = None
    finished = False
    counted_paragraphs = []
    display_captions = {
        f"{number}. {caption}"
        for number, _, caption in [*MAIN_FIGURES, *MAIN_TABLES]
    }
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        heading_match = re.fullmatch(r"Heading (\d+)", paragraph.style.name)
        heading_level = int(heading_match.group(1)) if heading_match else None
        if heading_level is not None and text == "Abstract":
            started = True
        if (
            acknowledgments_heading_level is not None
            and heading_level is not None
            and heading_level <= acknowledgments_heading_level
            and text != "Acknowledgments"
        ):
            finished = True
            break
        if not started or not text:
            continue
        if heading_level is not None and text == "Acknowledgments":
            acknowledgments_heading_level = heading_level
        if text in display_captions:
            continue
        counted_paragraphs.append(text)
    if acknowledgments_heading_level is not None and not finished:
        finished = True
    if not started or acknowledgments_heading_level is None or not finished:
        raise ValueError(
            f"{path.name}: cannot locate Abstract-through-Acknowledgments range"
        )
    return _count_words("\n".join(counted_paragraphs))


def _validate_manuscript_word_limit(path: Path, limit: int = 7000) -> int:
    word_count = _manuscript_word_count(path)
    if word_count > limit:
        raise ValueError(
            f"Manuscript is {word_count} words from Abstract through "
            f"Acknowledgments; limit is {limit}"
        )
    return word_count


def build_submission_documents(root: Path) -> dict[str, Path]:
    """Build and validate manuscript and submission-support documents."""
    summary = pd.read_csv(root / "results/policy_summary.csv")
    information = pd.read_csv(root / "results/information_value.csv")
    coverage = pd.read_csv(root / "data/processed/gbif_coverage_summary.csv")
    literature = pd.read_csv(root / "data/processed/bibliography.csv")
    checkpoint = json.loads((root / "results/checkpoints/pipeline.json").read_text())
    outputs = {
        "manuscript": _main_manuscript(
            root, summary, information, coverage, literature, checkpoint
        ),
        "tables": _editable_tables(root),
        "figure_deck": _figure_deck(root),
        "cover_page": _cover_page(root),
        "cover_letter": _cover_letter(root),
    }
    _validate_literature(outputs["manuscript"], literature)
    _validate_references(
        outputs["manuscript"],
        [spec[0] for spec in MAIN_FIGURES],
        [spec[0] for spec in MAIN_TABLES],
    )
    for stale in [
        "SUPPORTING_INFORMATION.md",
        "SUPPORTING_INFORMATION_INLINE.docx",
    ]:
        (root / "manuscript" / stale).unlink(missing_ok=True)
    for name, path in outputs.items():
        _validate_office_file(path)
        if name not in {"cover_page", "cover_letter"}:
            _scan_anonymous_office_file(path)
    validation = root / "manuscript/DOCUMENT_VALIDATION.txt"
    manuscript_text = _document_text(outputs["manuscript"])
    manuscript_word_count = _validate_manuscript_word_limit(outputs["manuscript"])
    abstract_text = (
        manuscript_text.split("Abstract", maxsplit=1)[1]
        .split("Keywords:", maxsplit=1)[0]
        .strip()
    )
    abstract_word_count = _count_words(abstract_text)
    keywords_text = (
        manuscript_text.split("Keywords:", maxsplit=1)[1]
        .split("Article impact statement:", maxsplit=1)[0]
        .strip()
    )
    keyword_count = len(
        [keyword for keyword in keywords_text.split(";") if keyword.strip()]
    )
    impact_statement = (
        manuscript_text.split("Article impact statement:", maxsplit=1)[1]
        .split("Introduction", maxsplit=1)[0]
        .strip()
    )
    if abstract_word_count > 300:
        raise ValueError(f"Abstract is {abstract_word_count} words")
    if not 5 <= keyword_count <= 8:
        raise ValueError(f"Keyword count is {keyword_count}")
    if len(impact_statement) > 140:
        raise ValueError(f"Impact statement is {len(impact_statement)} characters")
    validation.write_text(
        "\n".join(
            [
                "Office document validation: PASS",
                f"manuscript_sha256={hashlib.sha256(outputs['manuscript'].read_bytes()).hexdigest()}",
                f"editable_tables_sha256={hashlib.sha256(outputs['tables'].read_bytes()).hexdigest()}",
                f"figure_deck_sha256={hashlib.sha256(outputs['figure_deck'].read_bytes()).hexdigest()}",
                "manuscript_word_limit=7000",
                "manuscript_word_range=Abstract through Acknowledgments; "
                "display captions and table contents excluded",
                f"manuscript_abstract_through_acknowledgments_words={manuscript_word_count}",
                f"abstract_words={abstract_word_count}",
                f"keywords={keyword_count}",
                f"impact_statement_characters={len(impact_statement)}",
                f"main_figures={len(MAIN_FIGURES)}",
                f"main_tables={len(MAIN_TABLES)}",
                "supporting_information=none; all content integrated in main text",
                "cross_references=PASS",
                "anonymous_office_xml_scan=PASS",
                "zip_integrity=PASS",
            ]
        )
        + "\n",
        encoding="utf-8",
    )
    outputs["validation"] = validation
    return outputs
